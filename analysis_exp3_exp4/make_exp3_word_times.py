"""Experiment 3 — time taken to select each word, as an Excel workbook.

Sheets
  Every selection  one row per selection, in order, with its time
  By word          each vocabulary word, how often and how long on average
  By position      the 4x3 grid position map
  By run           one row per run
  Summary          headline numbers
"""
import csv, os, glob
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/advait/Downloads/EXP 3"
OUT  = "/Users/advait/Desktop/GazeTracking/Mobile_Gaze_tracking-main/analysis_exp3_exp4"
XL   = os.path.join(OUT, "exp3_word_times.xlsx")
S = sorted(glob.glob(os.path.join(ROOT, "session_*")))
R, C = 4, 3

def sect(p):
    blocks, cur = [], None
    for line in open(p):
        line = line.rstrip("\n")
        if not line.strip() or line.startswith("#"):
            cur = None; continue
        if cur is None:
            cur = {"h": line.split(","), "r": []}; blocks.append(cur)
        else:
            cur["r"].append(next(csv.reader([line])))
    return [[dict(zip(b["h"], r)) for r in b["r"]] for b in blocks]

sel_rows, run_rows = [], []
for s in S:
    sess = os.path.basename(s).replace("session_", "")
    sel, geom, ov = sect(os.path.join(s, "exp3/run1_communication/trials.csv"))
    ov = ov[0]
    target = {int(g["cell_idx"]): g["is_target_word"] == "1" for g in geom}
    for i, r in enumerate(sel, 1):
        k = int(r["cell_idx"])
        secs = float(r["since_prev_s"])
        sel_rows.append({
            "Session": sess,
            "Selection_no": i,
            "Word": r["word"],
            "Advanced_sentence": "yes" if r["correct"] == "1" else "no",
            "Sentence_position": int(r["target_pos"]) + 1 if int(r["target_pos"]) >= 0 else None,
            "Cell": k + 1,
            "Grid_row": k // C + 1,
            "Grid_col": k % C + 1,
            "In_target_sentence": "yes" if target[k] else "no",
            "Seconds_to_select": round(secs, 3),
            "Milliseconds": round(secs * 1000),
            "Elapsed_in_run_s": round(float(r["at_s"]), 3),
        })
    run_rows.append({
        "Session": sess,
        "Target_sentence": ov["target_sentence"],
        "Composed_sentence": ov["composed_sentence"],
        "Completed": "yes" if ov["completed"] == "1" else "no",
        "Selections": int(ov["n_selections"]),
        "Advanced": int(ov["correct"]),
        "Did_not_advance": int(ov["incorrect"]),
        "Selection_accuracy_pct": float(ov["selection_accuracy_pct"]),
        "Run_time_s": float(ov["duration_s"]),
        "Seconds_per_selection": float(ov["mean_s_per_selection"]),
        "Words_per_min": float(ov["words_per_min"]),
    })

sel_df = pd.DataFrame(sel_rows)
run_df = pd.DataFrame(run_rows)

by_word = (sel_df.groupby("Word")
           .agg(Times_selected=("Seconds_to_select", "size"),
                Mean_s=("Seconds_to_select", "mean"),
                Median_s=("Seconds_to_select", "median"),
                Fastest_s=("Seconds_to_select", "min"),
                Slowest_s=("Seconds_to_select", "max"),
                Distinct_cells_used=("Cell", "nunique"))
           .round(3).sort_values(["Times_selected", "Mean_s"], ascending=[False, True])
           .reset_index())

by_pos = (sel_df.groupby(["Grid_row", "Grid_col"])
          .agg(Selections=("Seconds_to_select", "size"),
               Mean_s=("Seconds_to_select", "mean"),
               Median_s=("Seconds_to_select", "median"),
               Fastest_s=("Seconds_to_select", "min"),
               Slowest_s=("Seconds_to_select", "max"))
          .round(3).reset_index())
by_pos["Cell"] = (by_pos["Grid_row"] - 1) * C + by_pos["Grid_col"]
for r in range(1, R + 1):
    for c in range(1, C + 1):
        if not ((by_pos["Grid_row"] == r) & (by_pos["Grid_col"] == c)).any():
            by_pos.loc[len(by_pos)] = {"Grid_row": r, "Grid_col": c, "Selections": 0,
                                       "Mean_s": None, "Median_s": None,
                                       "Fastest_s": None, "Slowest_s": None,
                                       "Cell": (r - 1) * C + c}
by_pos = by_pos.sort_values("Cell")[["Cell", "Grid_row", "Grid_col", "Selections",
                                     "Mean_s", "Median_s", "Fastest_s", "Slowest_s"]]

adv = sel_df[sel_df["Advanced_sentence"] == "yes"]
summary = pd.DataFrame([
    ("Sessions / runs", len(S)),
    ("Selections", len(sel_df)),
    ("Selections that advanced the sentence", len(adv)),
    ("Selection accuracy (%)", round(100 * len(adv) / len(sel_df), 2)),
    ("Runs completing the sentence", int((run_df["Completed"] == "yes").sum())),
    ("Distinct words selected", sel_df["Word"].nunique()),
    ("Mean seconds per selection", round(sel_df["Seconds_to_select"].mean(), 3)),
    ("Median seconds per selection", round(sel_df["Seconds_to_select"].median(), 3)),
    ("Fastest single selection (s)", round(sel_df["Seconds_to_select"].min(), 3)),
    ("Slowest single selection (s)", round(sel_df["Seconds_to_select"].max(), 3)),
    ("Mean run time (s)", round(run_df["Run_time_s"].mean(), 3)),
    ("Mean words per minute", round(run_df["Words_per_min"].mean(), 2)),
    ("Grid", f"{R} rows x {C} cols, 12 words, reshuffled each session"),
    ("Dwell threshold (s)", 1.0),
], columns=["Measure", "Value"])

with pd.ExcelWriter(XL, engine="openpyxl") as w:
    sel_df.to_excel(w, sheet_name="Every selection", index=False)
    by_word.to_excel(w, sheet_name="By word", index=False)
    by_pos.to_excel(w, sheet_name="By position", index=False)
    run_df.to_excel(w, sheet_name="By run", index=False)
    summary.to_excel(w, sheet_name="Summary", index=False)

from openpyxl import load_workbook
wb = load_workbook(XL)
HDR_FILL = PatternFill("solid", fgColor="12161B")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
BAND = PatternFill("solid", fgColor="F2F5F8")
thin = Side(style="thin", color="D6DBE0")
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        L = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(L + 3, 11), 44)
    for i, row in enumerate(ws.iter_rows(min_row=2)):
        for c in row:
            c.border = Border(bottom=thin)
            if isinstance(c.value, float): c.number_format = "0.000"
            if i % 2: c.fill = BAND
    for c in ws[1]:
        if str(c.value) in ("Target_sentence", "Composed_sentence"):
            ws.column_dimensions[c.column_letter].width = 34
wb.save(XL)

print(f"wrote {XL}")
print(f"  {len(sel_df)} selections over {len(S)} runs, {len(adv)} advanced the sentence")
print(f"  mean {sel_df['Seconds_to_select'].mean():.2f} s per selection "
      f"(median {sel_df['Seconds_to_select'].median():.2f}, "
      f"range {sel_df['Seconds_to_select'].min():.2f}–{sel_df['Seconds_to_select'].max():.2f})")
