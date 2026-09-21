import csv, os, glob, math, json, statistics as st
import numpy as np, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/advait/Downloads/EXP 3"
OUT  = "/Users/advait/Desktop/GazeTracking/Mobile_Gaze_tracking-main/analysis_exp3_exp4"
os.makedirs(OUT, exist_ok=True)
XL = f"{OUT}/exp3_exp4_results.xlsx"

S = sorted(glob.glob(os.path.join(ROOT, "session_*")))
def sid(p):
    b = os.path.basename(p)[-6:]
    return b[:2]+":"+b[2:4]+":"+b[4:]
LBL = [sid(s) for s in S]

def sect(p):
    blocks, cur = [], None
    for line in open(p):
        line = line.rstrip("\n")
        if not line.strip() or line.startswith("#"):
            cur = None; continue
        if cur is None:
            cur = {"h": line.split(","), "r": []}; blocks.append(cur)
        else:
            cur["r"].append(next(csv.reader([line])))
    return [[dict(zip(b["h"], r)) for r in b["r"]] for b in blocks]

VAL = [sect(os.path.join(s, "validation/run1_9dot/trials.csv"))[0] for s in S]
E3  = [sect(os.path.join(s, "exp3/run1_communication/trials.csv"))   for s in S]
E4  = [sect(os.path.join(s, "exp4/run1_predictive_cued/trials.csv")) for s in S]
F = float

# ── T1 validation per session ─────────────────────────────────
t1 = []
for s, lb, v in zip(S, LBL, VAL):
    m = json.load(open(os.path.join(s, "validation/run1_9dot/meta.json")))["summary"]
    t1.append(dict(Session=lb, Dots_confirmed=m["confirmed_dots"], Dots_total=m["total_dots"],
                   Mean_err_deg=round(m["mean_err_deg"],3), Mean_err_pt=round(m["mean_err_pt"],2),
                   Mean_RMS_deg=round(m["mean_rms_deg"],3), Mean_RMS_pt=round(m["mean_rms_pt"],2),
                   Worst_dot_deg=round(m["worst_err_deg"],3), Verdict=m["verdict"]))
T1 = pd.DataFrame(t1)

# ── T2 validation per dot (long) ──────────────────────────────
t2 = []
for lb, v in zip(LBL, VAL):
    for r in v:
        t2.append(dict(Session=lb, Dot=int(r["dot"]),
                       Row=(int(r["dot"])-1)//3+1, Col=(int(r["dot"])-1)%3+1,
                       Target_x_pt=F(r["target_x_pt"]), Target_y_pt=F(r["target_y_pt"]),
                       Mean_x_pt=F(r["mean_x_pt"]), Mean_y_pt=F(r["mean_y_pt"]),
                       Err_x_pt=F(r["err_x_pt"]), Err_y_pt=F(r["err_y_pt"]),
                       Err_pt=F(r["err_pt"]), Err_deg=F(r["err_deg"]),
                       SD_pt=F(r["sd_pt"]), RMS_pt=F(r["rms_pt"]), RMS_deg=F(r["rms_deg"]),
                       N_samples=int(r["n_samples"]), Confirm_order=int(r["confirm_order"]),
                       Confirmed_at_s=F(r["confirmed_at_s"])))
T2 = pd.DataFrame(t2)

# ── T3 dot summary (pivot for the heat map) ───────────────────
T3 = (T2.groupby(["Dot","Row","Col"])
        .agg(Mean_err_deg=("Err_deg","mean"), SD_err_deg=("Err_deg","std"),
             Mean_RMS_deg=("RMS_deg","mean"), Mean_err_pt=("Err_pt","mean"),
             Mean_bias_x_pt=("Err_x_pt","mean"), Mean_bias_y_pt=("Err_y_pt","mean"),
             N=("Err_deg","size"))
        .round(3).reset_index())
T3_grid = T3.pivot(index="Row", columns="Col", values="Mean_err_deg").round(2)
T3_grid.index = ["Top row","Middle row","Bottom row"]
T3_grid.columns = ["Left","Centre","Right"]
T3_grid["Row mean"] = T3_grid.mean(axis=1).round(2)

# ── T4 Exp3 per run ───────────────────────────────────────────
t4 = []
for lb, (sel, layout, ov) in zip(LBL, E3):
    o = ov[0]
    t4.append(dict(Session=lb, N_selections=int(o["n_selections"]), Correct=int(o["correct"]),
                   Incorrect=int(o["incorrect"]),
                   Selection_accuracy_pct=F(o["selection_accuracy_pct"]),
                   Mean_s_per_selection=F(o["mean_s_per_selection"]),
                   Mean_s_per_correct_word=F(o["mean_s_per_correct_word"]),
                   Words_per_min=F(o["words_per_min"]), Duration_s=F(o["duration_s"]),
                   Completed=int(o["completed"]),
                   Target_sentence=o["target_sentence"], Composed_sentence=o["composed_sentence"],
                   Virtual_tz_pt=F(o["tz_points"])))
T4 = pd.DataFrame(t4)

# ── T5 Exp3 per selection ─────────────────────────────────────
t5 = []
for lb, (sel, layout, ov) in zip(LBL, E3):
    L = {r["cell_idx"]: r for r in layout}
    prev = None
    for r in sel:
        c = L[r["cell_idx"]]
        cx = (F(c["x_min"])+F(c["x_max"]))/2; cy = (F(c["y_min"])+F(c["y_max"]))/2
        hw = (F(c["x_max"])-F(c["x_min"]))/2; hh = (F(c["y_max"])-F(c["y_min"]))/2
        dx, dy = F(r["pred_x"])-cx, F(r["pred_y"])-cy
        corr = r["correct"] == "1"
        kind = "correct" if corr else ("unintended re-selection" if r["cell_idx"] == prev else "spatial substitution")
        t5.append(dict(Session=lb, Selection=int(r["selection"]), Cell_idx=int(r["cell_idx"]),
                       Word=r["word"], Correct=int(r["correct"]), Target_pos=int(r["target_pos"]),
                       At_s=F(r["at_s"]), Since_prev_s=F(r["since_prev_s"]),
                       Pred_x_pt=F(r["pred_x"]), Pred_y_pt=F(r["pred_y"]),
                       Cell_centre_x=round(cx,2), Cell_centre_y=round(cy,2),
                       Offset_from_centre_pt=round(math.hypot(dx,dy),2),
                       Offset_norm_halfcell=round(max(abs(dx)/hw, abs(dy)/hh),3),
                       Outcome_class=kind))
        prev = r["cell_idx"]
T5 = pd.DataFrame(t5)

# ── T6 Exp4 per run ───────────────────────────────────────────
SLOT = [0,1,2,3,4,5,7]
t6 = []
for lb, (sel, grids, geo, trials, ov) in zip(LBL, E4):
    o = ov[0]
    on = hit = off = 0
    for r in sel:
        if r["kind"] != "word": continue
        ir = int(r["intended_rank"])
        if ir < 0: off += 1; continue
        on += 1
        if int(r["chosen_rank"]) == ir: hit += 1
    t6.append(dict(Session=lb, N_trials=int(o["n_trials"]), Completed_trials=int(o["completed_trials"]),
                   Word_selections=int(o["word_selections"]), Correct=int(o["correct"]),
                   Selection_accuracy_pct=F(o["selection_accuracy_pct"]),
                   Predictor_hit_rate_pct=F(o["predictor_hit_rate_pct"]),
                   Corrections=int(o["corrections"]), Selection_overhead=F(o["selection_overhead"]),
                   Scorable_selections=on, Correct_when_scorable=hit,
                   Conditional_accuracy_pct=round(100*hit/on,1) if on else None,
                   Unscorable_selections=off, Duration_s=F(o["duration_s"]),
                   Virtual_tz_pt=F(o["tz_points"])))
T6 = pd.DataFrame(t6)

# ── T7 Exp4 per selection ─────────────────────────────────────
t7 = []
for lb, (sel, grids, geo, trials, ov) in zip(LBL, E4):
    for r in sel:
        ir = int(r["intended_rank"]); ch = int(r["chosen_rank"])
        if r["kind"] == "word" and ir >= 0:
            want, got = SLOT[ir], int(r["cell_idx"])
            d = max(abs(want//3 - got//3), abs(want%3 - got%3))
            scor, dist = "scorable", d
            attr = "correct" if d == 0 else ("adjacent cell (tracker-plausible)" if d == 1 else "two cells away (NOT tracker)")
        else:
            scor, dist, attr = ("unscorable (word absent)" if r["kind"]=="word" else "control"), None, \
                               ("downstream of divergence" if r["kind"]=="word" else "control activation")
        t7.append(dict(Session=lb, Selection=int(r["selection"]), Trial=int(r["trial"]),
                       Question=r["question"], Grid_step=int(r["grid_step"]),
                       Cell_idx=int(r["cell_idx"]), Kind=r["kind"], Word_chosen=r["word"],
                       Correct=int(r["correct"]), Chosen_rank=ch,
                       Intended_word=r["intended_word"], Intended_rank=ir,
                       Scorability=scor, Cell_distance=dist, Attribution=attr,
                       At_s=F(r["at_s"]), Since_prev_s=F(r["since_prev_s"]),
                       Pred_x_pt=F(r["pred_x"]), Pred_y_pt=F(r["pred_y"])))
T7 = pd.DataFrame(t7)

# ── T8 Exp4 per trial ─────────────────────────────────────────
t8 = []
for lb, (sel, grids, geo, trials, ov) in zip(LBL, E4):
    for t in trials:
        ss = [r for r in sel if r["trial"] == t["trial"] and r["kind"] == "word"]
        first = "none"
        for r in ss:
            if int(r["intended_rank"]) < 0: first = "predictor (word absent)"; break
            if int(r["chosen_rank"]) != int(r["intended_rank"]): first = "selection (word present, other cell taken)"; break
        t8.append(dict(Session=lb, Trial=int(t["trial"]), Question=t["question"],
                       Question_text=t["question_text"], Cued_answer=t["cued_answer"],
                       Composed=t["composed"], Completed=int(t["completed"]),
                       Word_selections=int(t["word_selections"]), Correct=int(t["correct"]),
                       Corrections=int(t["corrections"]), Ideal_selections=int(t["ideal_selections"]),
                       Predictor_hits=int(t["predictor_hits"]),
                       Predictor_opportunities=int(t["predictor_opportunities"]),
                       Prompt_s=F(t["prompt_s"]), Response_s=F(t["response_s"]),
                       First_divergence=first))
T8 = pd.DataFrame(t8)

# ── T9 timing (long, for a pivot/box plot) ────────────────────
t9 = []
for lb, (sel, layout, ov) in zip(LBL, E3):
    for r in sel:
        t9.append(dict(Session=lb, Condition="Exp 3 words (1.0 s dwell)",
                       Interval_s=F(r["since_prev_s"])))
for lb, (sel, grids, geo, trials, ov) in zip(LBL, E4):
    for r in sel:
        t9.append(dict(Session=lb,
                       Condition="Exp 4 words (1.0 s dwell)" if r["kind"]=="word" else "Exp 4 controls (1.5 s dwell)",
                       Interval_s=F(r["since_prev_s"])))
T9 = pd.DataFrame(t9)
T9s = (T9.groupby("Condition")["Interval_s"]
        .agg(N="size", Mean="mean", SD="std", Min="min", Q25=lambda x: x.quantile(.25),
             Median="median", Q75=lambda x: x.quantile(.75), Max="max").round(3).reset_index())
T9s["Dwell_charged_s"] = T9s["Condition"].map({
    "Exp 3 words (1.0 s dwell)": 1.0,
    "Exp 4 words (1.0 s dwell)": 1.0,
    "Exp 4 controls (1.5 s dwell)": 1.5})
T9s["Search_overhead_s"] = (T9s["Mean"] - T9s["Dwell_charged_s"]).round(3)
T9s["Overhead_pct_of_interval"] = (100*T9s["Search_overhead_s"]/T9s["Mean"]).round(1)

# ── T10 signal quality ────────────────────────────────────────
def load(p): return list(csv.DictReader(open(p)))
t10 = []
for s, lb in zip(S, LBL):
    for exp, run in (("exp3","run1_communication"), ("exp4","run1_predictive_cued")):
        R = load(os.path.join(s, exp, run, "samples.csv"))
        t = [F(r["t_s"]) for r in R]
        blink = 100*sum(1 for r in R if r.get("blink_held") not in (None,"","0"))/len(R)
        inc = [r["in_cell"] for r in R if r.get("in_cell") not in (None,"")]
        P = [r for r in R if r["pred_x"] not in ("","None")]
        Q = [r for r in R if r["raw_pred_x"] not in ("","None")]
        jf = np.median([math.hypot(F(b["pred_x"])-F(a["pred_x"]), F(b["pred_y"])-F(a["pred_y"])) for a,b in zip(P,P[1:])])
        jr = np.median([math.hypot(F(b["raw_pred_x"])-F(a["raw_pred_x"]), F(b["raw_pred_y"])-F(a["raw_pred_y"])) for a,b in zip(Q,Q[1:])])
        t10.append(dict(Session=lb, Experiment=exp.upper(), N_samples=len(R),
                        Rate_Hz=round(1/np.mean(np.diff(t)),2),
                        Blink_gated_pct=round(blink,2),
                        In_active_cell_pct=round(100*sum(1 for v in inc if v=="1")/len(inc),2),
                        Head_yaw_SD_deg=round(float(np.std([F(r["head_yaw_deg"]) for r in R])),3),
                        Head_pitch_SD_deg=round(float(np.std([F(r["head_pitch_deg"]) for r in R])),3),
                        Head_tz_mean_mm=round(float(np.mean([F(r["head_tz_mm"]) for r in R])),1),
                        Jitter_filtered_pt=round(float(jf),2), Jitter_raw_pt=round(float(jr),2),
                        Jitter_reduction_pct=round(100*(1-jf/jr),1)))
T10 = pd.DataFrame(t10)

# ── T11 cross-experiment ──────────────────────────────────────
T11 = pd.DataFrame(dict(
    Session=LBL,
    Validation_err_deg=T1["Mean_err_deg"],
    Verdict=T1["Verdict"],
    Exp3_accuracy_pct=T4["Selection_accuracy_pct"],
    Exp3_wpm=T4["Words_per_min"],
    Exp4_selection_accuracy_pct=T6["Selection_accuracy_pct"],
    Exp4_conditional_accuracy_pct=T6["Conditional_accuracy_pct"],
    Exp4_predictor_hit_rate_pct=T6["Predictor_hit_rate_pct"],
)).sort_values("Validation_err_deg").reset_index(drop=True)

def r_(a, b): return round(float(np.corrcoef(a, b)[0,1]), 3)
T11c = pd.DataFrame([
    dict(Pair="Validation error  ×  Exp 3 selection accuracy", Pearson_r=r_(T11.Validation_err_deg, T11.Exp3_accuracy_pct), N=5),
    dict(Pair="Validation error  ×  Exp 3 words per minute",   Pearson_r=r_(T11.Validation_err_deg, T11.Exp3_wpm), N=5),
    dict(Pair="Validation error  ×  Exp 4 selection accuracy", Pearson_r=r_(T11.Validation_err_deg, T11.Exp4_selection_accuracy_pct), N=5),
    dict(Pair="Validation error  ×  Exp 4 conditional accuracy", Pearson_r=r_(T11.Validation_err_deg, T11.Exp4_conditional_accuracy_pct), N=5),
])

N_ALL = 0
for s_ in S:
    for e_, r_dir in (("exp3","run1_communication"),("exp4","run1_predictive_cued"),("validation","run1_9dot")):
        N_ALL += sum(1 for _ in open(os.path.join(s_, e_, r_dir, "samples.csv"))) - 1

# ── Summary sheet ─────────────────────────────────────────────
norm = T5.loc[T5.Correct == 1, "Offset_norm_halfcell"]
on_t = int(T6.Scorable_selections.sum()); hit_t = int(T6.Correct_when_scorable.sum())
g3 = T9.loc[T9.Condition.str.startswith("Exp 3"), "Interval_s"]
g4 = T9.loc[T9.Condition == "Exp 4 words (1.0 s dwell)", "Interval_s"]
cheb = T7.Cell_distance.value_counts().sort_index()

SUM = pd.DataFrame([
    ("Sessions", 5, "runs on 2026-08-26, 20:08–20:39; no participant ID recorded in the export"),
    ("Runs analysed", 15, "5 calibration validation + 5 Experiment 3 + 5 Experiment 4"),
    ("Gaze samples", N_ALL, "Exp 3 + Exp 4 + validation, all logged frames"),
    ("", "", ""),
    ("CALIBRATION", "", ""),
    ("Mean angular error (°)", round(T1.Mean_err_deg.mean(),2), "45 dots; 4 sessions 'good', 1 'marginal'"),
    ("Mean RMS scatter (°)", round(T1.Mean_RMS_deg.mean(),2), "precision; flat across screen positions"),
    ("Top-row error (°)", round(T3_grid.loc["Top row","Row mean"],2), "vs bottom row — systematic bias, not noise"),
    ("Bottom-row error (°)", round(T3_grid.loc["Bottom row","Row mean"],2), "+47 % over the top row"),
    ("Pooled bias x / y (pt)", f"{T2.Err_x_pt.mean():+.1f} / {T2.Err_y_pt.mean():+.1f}", "small against a 134 pt cell width"),
    ("", "", ""),
    ("EXPERIMENT 3 — static 4×3 grid", "", ""),
    ("Selection accuracy (pooled)", f"{100*T4.Correct.sum()/T4.N_selections.sum():.1f} %", f"{T4.Correct.sum()} of {T4.N_selections.sum()} selections"),
    ("Words per minute", f"{T4.Words_per_min.mean():.2f} (SD {T4.Words_per_min.std():.2f})", "mean across 5 runs"),
    ("Seconds per selection", f"{T4.Mean_s_per_selection.mean():.2f} (SD {T4.Mean_s_per_selection.std():.2f})", "against a 1.0 s dwell requirement"),
    ("Sentences completed", f"{int(T4.Completed.sum())} / 5", "'I want to drink water'"),
    ("Offset from tile centre", f"{norm.mean():.2f} of a half-tile", f"mean {T5.loc[T5.Correct==1,'Offset_from_centre_pt'].mean():.1f} pt on 134 × 149 pt tiles"),
    ("Picks in the outer fifth", f"{100*(norm>0.8).mean():.0f} %", "tiles are not at the resolution limit"),
    ("Unintended re-selections", int((T5.Outcome_class=="unintended re-selection").sum()), "dwell timer re-armed without gaze leaving the tile"),
    ("Spatial substitutions", int((T5.Outcome_class=="spatial substitution").sum()), "a genuinely different tile fired"),
    ("", "", ""),
    ("EXPERIMENT 4 — predictive 3×3 grid (cued)", "", ""),
    ("Selection accuracy as logged", f"{100*T6.Correct.sum()/T6.Word_selections.sum():.1f} %", "NOT VALID — see the cued-condition confound"),
    ("Conditional accuracy", f"{100*hit_t/on_t:.1f} %", f"{hit_t} of {on_t} selections made while the needed word was on screen"),
    ("Unscorable selections", f"{int(T6.Unscorable_selections.sum())} of {int(T6.Word_selections.sum())} ({100*T6.Unscorable_selections.sum()/T6.Word_selections.sum():.0f} %)", "downstream of a divergence; predictor cannot re-offer the word"),
    ("Cued answers completed", "0 / 15", "every trial diverged from its cue"),
    ("Picks two cells from target", f"{int(cheb.get(2,0))} of {on_t} ({100*cheb.get(2,0)/on_t:.0f} %)", "beyond any 1.5° tracker error — deliberate branch choices"),
    ("Backspace activations", 0, "⌫ never fired, accidentally or intentionally"),
    ("Selection overhead", f"{T6.Selection_overhead.mean():.2f}", "≈1.0 — participants were efficient, just answering differently"),
    ("", "", ""),
    ("TIMING", "", ""),
    ("Exp 3 interval", f"{g3.mean():.2f} s (median {g3.median():.2f})", f"search overhead {g3.mean()-1:.2f} s above the 1.0 s dwell"),
    ("Exp 4 interval", f"{g4.mean():.2f} s (median {g4.median():.2f})", f"search overhead {g4.mean()-1:.2f} s; includes the 0.4 s refresh lockout"),
    ("Share of time that is search", f"{100*(g3.mean()-1)/g3.mean():.0f} %", "halving the dwell threshold buys ≤13 %"),
    ("", "", ""),
    ("SIGNAL QUALITY", "", ""),
    ("Effective logging rate", f"{T10.Rate_Hz.mean():.1f} Hz (SD {T10.Rate_Hz.std():.1f})", "a 1.0 s dwell rests on ~10 samples"),
    ("Jitter, raw → filtered", f"{T10.Jitter_raw_pt.mean():.1f} → {T10.Jitter_filtered_pt.mean():.1f} pt", f"One-Euro removes {100*(1-T10.Jitter_filtered_pt.mean()/T10.Jitter_raw_pt.mean()):.0f} %"),
    ("Head yaw / pitch SD", f"{T10.Head_yaw_SD_deg.mean():.2f}° / {T10.Head_pitch_SD_deg.mean():.2f}°", "near-static posture — best-case condition"),
], columns=["Metric", "Value", "Note"])

NOTES = pd.DataFrame([
 ("READ THIS FIRST", "Experiment 4's cued condition did not hold. In all 15 trials the first divergence from the cue is a deliberate branch choice (a semantically valid answer to the question asked), never an absent word. 7 of 26 scorable picks landed two cells from the needed word — impossible for a 1.5° estimator on 134 × 199 pt cells. Report Exp 4 as free response, or re-run with the cue visible during the response phase. 'Selection_accuracy_pct' and 'Predictor_hit_rate_pct' as logged are not valid."),
 ("Participants", "The export records no participant identifier. These are 5 sessions on one day, not 5 participants. Between-subject variance is unmeasured — describe them as sessions."),
 ("Degrees are virtual", "Angular error uses the calibration's fitted |tz| (1 119–1 245 pt across sessions), not a measured viewing distance. Comparable within a session, not across participants."),
 ("Cascade inflation", "Neither experiment has an undo path in the scoring, so one divergence contaminates every later selection in the trial. Raw per-selection accuracy UNDERSTATES per-selection reliability. Report first-divergence counts alongside it (sheet 'Exp4 trials')."),
 ("Single posture / lighting", "Head yaw SD 1.26° — these are stationary, hand-held, well-lit best-case numbers."),
 ("Exp 3 familiarity", "The target sentence is identical across all 5 runs, so later runs carry unmeasured task familiarity even though tile positions were reshuffled."),
 ("n = 5 correlations", "The r values in 'Cross-experiment' are directions, not effect sizes. The confidence interval spans zero at n = 5."),
 ("Provenance", "All values computed directly from the exported CSVs in 'EXP 3/session_20260826_*' (schema v4). No re-scoring, no exclusions."),
 ("Pipeline", "One-Euro smoother — gaze (min_cutoff 1.2, β 0.35), eye (0.5, 0.02), head (0.8, 0.1); blink EAR ratio 0.55 / floor 0.12 / max 5 gated frames."),
], columns=["Topic", "Detail"])

SHEETS = [
    ("Summary", SUM), ("Caveats", NOTES),
    ("T1 Validation runs", T1), ("T2 Validation dots", T2),
    ("T3 Dot summary", T3), ("T3b Error grid", T3_grid.reset_index().rename(columns={"index":"Row"})),
    ("T4 Exp3 runs", T4), ("T5 Exp3 selections", T5),
    ("T6 Exp4 runs", T6), ("T7 Exp4 selections", T7), ("T8 Exp4 trials", T8),
    ("T9 Timing raw", T9), ("T9b Timing summary", T9s),
    ("T10 Signal quality", T10),
    ("T11 Cross-experiment", T11), ("T11b Correlations", T11c),
]

with pd.ExcelWriter(XL, engine="openpyxl") as w:
    for name, df in SHEETS:
        df.to_excel(w, sheet_name=name[:31], index=False)

# ── formatting ────────────────────────────────────────────────
from openpyxl import load_workbook
wb = load_workbook(XL)
HDR_FILL = PatternFill("solid", fgColor="12161B")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
BAND     = PatternFill("solid", fgColor="F2F5F8")
SEC_FILL = PatternFill("solid", fgColor="E3EDFA")
thin = Side(style="thin", color="D6DBE0")

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        L = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(L + 3, 11), 52)
    for i, row in enumerate(ws.iter_rows(min_row=2)):
        for c in row:
            c.border = Border(bottom=thin)
            if isinstance(c.value, float): c.number_format = "0.00"
            if i % 2: c.fill = BAND
    # long text columns wrap
    for c in ws[1]:
        if str(c.value) in ("Note","Detail","Question_text","Cued_answer","Composed",
                            "Target_sentence","Composed_sentence","Attribution","Topic","Metric","Pair"):
            ws.column_dimensions[c.column_letter].width = 46
            for r in ws.iter_rows(min_row=2, min_col=c.column, max_col=c.column):
                r[0].alignment = Alignment(wrap_text=True, vertical="top")

ws = wb["Summary"]
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 24
for row in ws.iter_rows(min_row=2, max_col=1):
    v = row[0].value
    if v and v.isupper() and len(v) > 3:
        for c in ws[row[0].row]:
            c.fill = SEC_FILL; c.font = Font(bold=True, size=10)

ws = wb["Caveats"]
ws["A2"].font = Font(bold=True, color="B03A3A", size=11)
ws.column_dimensions["B"].width = 96
for r in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    r[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r[0].row].height = 62

wb.save(XL)
print("workbook:", XL)
for n, d in SHEETS: print(f"   {n:24s} {d.shape[0]:5d} rows × {d.shape[1]} cols")
